#!/usr/bin/env python3
"""Page Guide translation review: export a side-by-side sheet for human reviewers, then import their decisions.

  python3 scripts/translation_review.py export [--checks-el FILE.md] [--checks-es FILE.md] [--out review.xlsx]
  python3 scripts/translation_review.py import --sheet review.xlsx [--mark-reviewed el|es]

Export: one row per string (English source, Greek draft, Spanish draft, automatic-check notes, decision columns).
Import: applies "Edit" corrections to data/guide_i18n.json. A language is only switched to "reviewed" when the
director passes --mark-reviewed for it AND every row of that language is Approve or Edit AND the reviewer's name
and date are filled in on the sheet. The script never fills in a reviewer name. English is never changed.
After import: python3 build_guide_index.py && python3 build_guide.py && python3 scripts/check_guide.py
"""
import argparse, hashlib, json, os, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
I18N = os.path.join(ROOT, "data", "guide_i18n.json")
GUIDE = os.path.join(ROOT, "data", "guide.json")
LANGS = ("el", "es")
DECISIONS = ["Approve", "Edit", "Question"]
COLS = ["Segment", "English (source)", "Greek draft", "Greek automatic check", "Greek decision", "Greek correction",
        "Spanish draft", "Spanish automatic check", "Spanish decision", "Spanish correction", "Reviewer notes", "English hash"]


def rows():
    """Yield (segment_id, english, {lang: text}, setter) for every translatable string, ids matching the checkers."""
    d = json.load(open(I18N, encoding="utf-8"))
    g = json.load(open(GUIDE, encoding="utf-8"))
    out = []
    for k in d["ui"]["en"]:
        en_ui = d["ui"]["en"][k] if d["ui"]["en"][k] is not None else g.get(k, "")  # intro lives in guide.json
        out.append((f"ui.{k}", en_ui, {l: d["ui"][l].get(k, "") for l in LANGS}, ("ui", k, None)))
    for k, pair in d["labels"]["en"].items():
        for i, en in enumerate(pair):
            out.append((f"labels.{k}[{i}]", en, {l: d["labels"][l][k][i] for l in LANGS}, ("labels", k, i)))
    for k, en in g["status_text"].items():
        out.append((f"status.{k}", en, {l: d["status"][l].get(k, "") for l in LANGS}, ("status", k, None)))
    titles = {gr["id"]: gr["title"] for gr in g["groups"]}
    for k in d["groups"]["el"]:
        out.append((f"groups.{k}", titles.get(k, ""), {l: d["groups"][l].get(k, "") for l in LANGS}, ("groups", k, None)))
    en_items = {it["id"]: it for gr in g["groups"] for it in gr["items"]}
    for iid, tr in d["items"].items():
        en = en_items[iid]
        en_parts = [en["q"]] + list(en["a"])
        for i, en_text in enumerate(en_parts):
            out.append((f"items.{iid}[{i}]", en_text, {l: (tr[l][i] if i < len(tr[l]) else "") for l in LANGS}, ("items", iid, i)))
    return d, out


def parse_checks(path):
    notes = {}
    if path and os.path.exists(path):
        for line in open(path, encoding="utf-8"):
            m = re.match(r"\| (\S+) \| ([^|]+) \| ([^|]+) \| ([^|]*) \| ([^|]*) \|", line)
            if m and m.group(1) not in ("Segment", "---"):
                notes.setdefault(m.group(1), []).append(f"{m.group(3).strip()}: {m.group(4).strip()} ({m.group(5).strip()})")
    return notes


def h(s):
    return hashlib.sha256(s.encode("utf-8")).hexdigest()[:12]


def export(a):
    d, R = rows()
    ce, cs = parse_checks(a.checks_el), parse_checks(a.checks_es)
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(COLS)
    for sid, en, tr, _ in R:
        ws.append([sid, en, tr["el"], "; ".join(ce.get(sid, [])) or "none", "", "", tr["es"], "; ".join(cs.get(sid, [])) or "none", "", "", "", h(en)])
    head = Font(name="Arial", bold=True, color="FFFFFF")
    for c in ws[1]:
        c.font = head
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [24, 50, 50, 26, 12, 40, 50, 26, 12, 40, 30, 12]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for idx in (4, 5, 8, 9, 10):
            row[idx].fill = yellow
    n = ws.max_row
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:E{n}")
    dv.add(f"I2:I{n}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:L{n}"
    lg = wb.create_sheet("Read me")
    lines = [
        ["KRIOS Page Guide translation review", ""],
        ["What to do", "For each row, choose a decision in the yellow column for your language: Approve (the draft is right), Edit (type the corrected text in the correction column), or Question (explain in Reviewer notes)."],
        ["Edit only", "the yellow cells. Do not change the English, the drafts or the Segment column."],
        ["Automatic check", "Notes from the spelling and style checkers (first pass only). 'none' means nothing was flagged; it is not an approval."],
        ["House rules", "Greek: monotonic, « » quotes, ; for questions, no em dashes. Spanish: neutral Latin American, usted, ¿ ¡, no em dashes, numbers as 1200 and 430 000."],
        ["Labels", "HISTORY, TRADITION, BELIEF, PRACTICE must stay consistent. BELIEF stays framed as belief."],
        ["Example row", "items.name-day[0] | Greek decision: Approve | Spanish decision: Edit | Spanish correction: (corrected sentence)"],
        ["", ""],
        ["Greek reviewer name", ""],
        ["Greek review date", ""],
        ["Spanish reviewer name", ""],
        ["Spanish review date", ""],
        ["Exported", f"{len(R)} strings from data/guide_i18n.json"],
    ]
    for r in lines:
        lg.append(r)
    lg.column_dimensions["A"].width = 24
    lg.column_dimensions["B"].width = 110
    for row in lg.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10, bold=(c.column == 1))
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in (9, 10, 11, 12):
        lg.cell(r, 2).fill = yellow
    lg["A1"].font = Font(name="Arial", size=13, bold=True)
    out = a.out or os.path.join(ROOT, "docs", "guide", "Page_Guide_translation_review.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"export: {len(R)} strings -> {out}")


def import_(a):
    d, R = rows()
    wb = load_workbook(a.sheet)
    ws, lg = wb["Review"], wb["Read me"]
    meta = {lg.cell(r, 1).value: lg.cell(r, 2).value for r in range(1, lg.max_row + 1)}
    by_id = {r[0]: r for r in R}
    hdr = [c.value for c in ws[1]]
    col = {name: hdr.index(name) for name in COLS}
    stats = {l: {"Approve": 0, "Edit": 0, "Question": 0, "blank": 0} for l in LANGS}
    changed, problems = 0, []
    lname = {"el": "Greek", "es": "Spanish"}
    for row in ws.iter_rows(min_row=2, values_only=True):
        sid = row[col["Segment"]]
        if sid not in by_id:
            problems.append(f"unknown segment {sid}")
            continue
        _, en, tr, (sec, key, idx) = by_id[sid]
        if row[col["English hash"]] != h(en):
            problems.append(f"{sid}: English changed since export; re-export before importing")
            continue
        for l in LANGS:
            dec = (row[col[f"{lname[l]} decision"]] or "").strip()
            stats[l][dec if dec in DECISIONS else "blank"] += 1
            if dec == "Edit":
                new = (row[col[f"{lname[l]} correction"]] or "").strip()
                if not new:
                    problems.append(f"{sid} {l}: Edit with no correction")
                    continue
                if "—" in new:
                    problems.append(f"{sid} {l}: correction contains an em dash; not applied")
                    continue
                target = d[sec][l] if sec != "items" else d["items"][key][l]
                if sec == "items":
                    target[idx] = new
                elif sec == "labels":
                    target[key][idx] = new
                else:
                    target[key] = new
                changed += 1
    for l in (a.mark_reviewed or []):
        s = stats[l]
        who, when = meta.get(f"{lname[l]} reviewer name"), meta.get(f"{lname[l]} review date")
        if s["blank"] or s["Question"] or not who or not when:
            problems.append(f"{l}: not marked reviewed ({s['blank']} blank, {s['Question']} questions, reviewer name {'set' if who else 'missing'}, date {'set' if when else 'missing'})")
        else:
            d["review"][l] = "reviewed"
            d.setdefault("review_record", {})[l] = {"reviewer": str(who), "date": str(when), "sheet": os.path.basename(a.sheet)}
    json.dump(d, open(I18N, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"import: {changed} corrections applied; decisions {stats}; review status {d['review']}")
    for p in problems:
        print("  !", p)
    if problems:
        sys.exit(1)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    e = sp.add_parser("export"); e.add_argument("--checks-el"); e.add_argument("--checks-es"); e.add_argument("--out")
    i = sp.add_parser("import"); i.add_argument("--sheet", required=True); i.add_argument("--mark-reviewed", nargs="*", choices=LANGS)
    a = ap.parse_args()
    {"export": export, "import": import_}[a.cmd](a)
